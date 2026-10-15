#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
E2E-тест конвертера документов против ЖИВОГО движка (http://127.0.0.1:8765).

Проверяет сценарий пользователя:
  «конвертация из ворд в тхт происходит некорректно — какие-то крякозябры,
   и нет конвертации в пдф и назад»

Что делается:
  0. /api/health — версия >= 1.3.0, docs/jobs/queue = true
  1. DOCX (кириллица + таблица) -> TXT   — результат строго UTF-8, без крякозябр
  2. DOCX -> PDF                          — кириллица извлекается из PDF (pypdf)
  3. TXT в CP1251 -> PDF                  — файл «из Блокнота русской Windows»
  4. TXT в CP1251 -> DOCX                 — кириллица на месте (python-docx)
  5. TXT в CP1251 -> TXT                  — перекодировка в UTF-8 (фикс кодировки)
  6. PDF -> DOCX                          — «конвертация назад» (pdf2docx)
  7. PDF -> TXT                           — «конвертация назад» (pypdf)

Запускать питоном окружения движка (нужны python-docx и pypdf):
  python tests/e2e_doc_conversion.py
Переменные окружения:
  VPAI_ENGINE  — базовый URL движка (по умолчанию http://127.0.0.1:8765)
"""
import io
import json
import os
import sys
import tempfile
import time
import urllib.parse
import urllib.request

# Кириллица в консоли Windows (CP866/CP1251 по умолчанию)
try:
    sys.stdout.reconfigure(encoding="utf-8")
    sys.stderr.reconfigure(encoding="utf-8")
except Exception:
    pass

BASE = os.environ.get("VPAI_ENGINE", "http://127.0.0.1:8765").rstrip("/")
JOB_TIMEOUT = 240  # сек на задачу (pdf2docx бывает медленный)
POLL_INTERVAL = 1.0

# Текст с «трудными» для кодировок символами: ё, Щ, №, «кавычки», латиница
SAMPLE_LINES = [
    "Привет из VoxPaperAi — проверка конвертации",
    "Щёлково: ёлка, ёжик и чаща — буква ё и щ",
    "Строка №2 с «кавычками» и числом 2026",
]
TABLE_CELLS = ["Ячейка А1", "Вторая строка таблицы"]

failures = []
passed = 0


def report(name: str, ok: bool, detail: str = "") -> None:
    global passed
    if ok:
        passed += 1
        print(f"  [PASS] {name}")
    else:
        failures.append((name, detail))
        print(f"  [FAIL] {name}  --  {detail}")


# ───────────────────────── HTTP-хелперы ─────────────────────────

def http_get(path: str, timeout: float = 30.0) -> tuple[int, bytes]:
    req = urllib.request.Request(BASE + path, headers={"User-Agent": "vpai-e2e/1.0"})
    with urllib.request.urlopen(req, timeout=timeout) as resp:
        return resp.status, resp.read()


def http_post_json(path: str, payload: dict, timeout: float = 30.0) -> dict:
    data = json.dumps(payload).encode("utf-8")
    req = urllib.request.Request(
        BASE + path,
        data=data,
        headers={"Content-Type": "application/json", "User-Agent": "vpai-e2e/1.0"},
        method="POST",
    )
    with urllib.request.urlopen(req, timeout=timeout) as resp:
        return json.loads(resp.read().decode("utf-8"))


def run_convert_job(file_path: str, target: str, label: str) -> tuple[bytes, dict]:
    """Создаёт job конвертации, ждёт завершения, скачивает результат. Возвращает (байты, снапшот job)."""
    r = http_post_json(
        "/api/jobs",
        {
            "type": "convert",
            "params": {"file_path": file_path, "target_format": target},
            "label": label,
        },
    )
    if r.get("status") != "started" or not r.get("job_id"):
        raise RuntimeError(f"POST /api/jobs ответил {r}")
    jid = r["job_id"]
    deadline = time.time() + JOB_TIMEOUT
    snapshot = {}
    while time.time() < deadline:
        snapshot = json.loads(http_get(f"/api/jobs/{jid}")[1])
        st = snapshot.get("status")
        if st in ("done", "error", "cancelled"):
            break
        time.sleep(POLL_INTERVAL)
    else:
        raise TimeoutError(f"job {jid} не завершился за {JOB_TIMEOUT} с: {snapshot.get('progress', {}).get('stage')}")
    if snapshot.get("status") != "done":
        raise RuntimeError(
            f"job {jid} статус={snapshot.get('status')}, "
            f"ошибка={snapshot.get('error')!r}, лог={snapshot.get('progress', {}).get('log', [])[-3:]}"
        )
    file_url = (snapshot.get("result") or {}).get("file")
    if not file_url:
        raise RuntimeError(f"в result нет файла: {snapshot.get('result')}")
    code, blob = http_get(file_url, timeout=60.0)
    if code != 200 or not blob:
        raise RuntimeError(f"скачивание {file_url} → HTTP {code}, {len(blob)} байт")
    return blob, snapshot


# ───────────────────────── Создание входных файлов ─────────────────────────

def make_fixtures(tmp: str) -> dict[str, str]:
    from docx import Document

    files = {}

    # DOCX с кириллицей и таблицей (тот же OOXML-формат, что у MS Word)
    doc = Document()
    for line in SAMPLE_LINES:
        doc.add_paragraph(line)
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = TABLE_CELLS[0]
    table.cell(0, 1).text = "B2 latin"
    table.cell(1, 0).text = TABLE_CELLS[1]
    table.cell(1, 1).text = "2026"
    # Специально НЕ-ASCII имя файла — проверка путей с кириллицей на Windows
    docx_path = os.path.join(tmp, "документ тест.docx")
    doc.save(docx_path)
    files["docx"] = docx_path

    # TXT в CP1251 — классический источник «крякозябр» (Блокнот русской Windows)
    cp1251_text = "\n".join(SAMPLE_LINES) + "\n"
    txt_cp = os.path.join(tmp, "текст_cp1251.txt")
    with open(txt_cp, "wb") as f:
        f.write(cp1251_text.encode("cp1251"))  # намеренно НЕ utf-8
    files["txt_cp1251"] = txt_cp

    # TXT в UTF-8 — базовый случай
    txt_u8 = os.path.join(tmp, "текст_utf8.txt")
    with open(txt_u8, "w", encoding="utf-8") as f:
        f.write(cp1251_text)
    files["txt_utf8"] = txt_u8
    return files


# ───────────────────────── Проверки результатов ─────────────────────────

def extract_pdf_text(blob: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(blob))
    return "\n".join((page.extract_text() or "") for page in reader.pages)


def extract_docx_text(blob: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(blob))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def has_all_cyrillic(text: str) -> tuple[bool, str]:
    for needle in ["Щёлково", "ёлка", "ёжик", "чаща", "«кавычками»", "2026", "VoxPaperAi"]:
        if needle not in text:
            return False, f"нет фрагмента {needle!r}"
    return True, "все фрагменты на месте"


# ───────────────────────── Тесты ─────────────────────────

def test_health() -> None:
    print("Тест 0: /api/health — версия и флаги docs/jobs/queue")
    try:
        h = json.loads(http_get("/api/health")[1])
        report("health отвечает", True)
    except Exception as e:
        report("health отвечает", False, str(e))
        return
    ver = str(h.get("version", ""))
    try:
        v_ok = tuple(int(x) for x in ver.split(".")) >= (1, 3, 0)
    except ValueError:
        v_ok = False
    report(f"версия {ver or '(пусто)'} >= 1.3.0", v_ok)
    report(f"docs={h.get('docs')} (библиотеки конвертации)", h.get("docs") is True)
    report(f"jobs={h.get('jobs')}", h.get("jobs") is True)
    report(f"queue={h.get('queue')}", h.get("queue") is True)


def test_docx_to_txt(files: dict) -> None:
    print("Тест 1: DOCX (кириллица) → TXT — главный сценарий «крякозябр»")
    try:
        blob, _ = run_convert_job(files["docx"], "txt", "e2e docx→txt")
    except Exception as e:
        report("docx→txt задача", False, str(e))
        return
    try:
        text = blob.decode("utf-8")  # СТРОГО utf-8: любой мусор = крякозябры
        ok, why = has_all_cyrillic(text)
        report("docx→txt: строго UTF-8 + кириллица + таблица", ok, why)
        report(
            "docx→txt: текст таблицы извлечён",
            TABLE_CELLS[0] in text and TABLE_CELLS[1] in text,
            "ячейки таблицы не найдены в TXT",
        )
    except UnicodeDecodeError as e:
        report("docx→txt: строго UTF-8 + кириллица", False, f"файл НЕ в UTF-8 (крякозябры): {e}")


def test_docx_to_pdf(files: dict, artifacts: dict) -> None:
    print("Тест 2: DOCX → PDF (без MS Word, через reportlab)")
    try:
        blob, _ = run_convert_job(files["docx"], "pdf", "e2e docx→pdf")
    except Exception as e:
        report("docx→pdf задача", False, str(e))
        return
    pdf_path = os.path.join(artifacts["tmp"], "результат_docx.pdf")
    with open(pdf_path, "wb") as f:
        f.write(blob)
    artifacts["pdf"] = pdf_path
    try:
        text = extract_pdf_text(blob)
        ok, why = has_all_cyrillic(text)
        report("docx→pdf: PDF валиден, кириллица извлекается", ok, why)
    except Exception as e:
        report("docx→pdf: PDF валиден", False, f"pypdf не прочитал PDF: {e}")


def test_txt_cp1251_to_pdf(files: dict) -> None:
    print("Тест 3: TXT в CP1251 → PDF")
    try:
        blob, _ = run_convert_job(files["txt_cp1251"], "pdf", "e2e cp1251→pdf")
        ok, why = has_all_cyrillic(extract_pdf_text(blob))
        report("cp1251→pdf: кириллица без крякозябр", ok, why)
    except Exception as e:
        report("cp1251→pdf", False, str(e))


def test_txt_cp1251_to_docx(files: dict) -> None:
    print("Тест 4: TXT в CP1251 → DOCX")
    try:
        blob, _ = run_convert_job(files["txt_cp1251"], "docx", "e2e cp1251→docx")
        ok, why = has_all_cyrillic(extract_docx_text(blob))
        report("cp1251→docx: кириллица на месте", ok, why)
    except Exception as e:
        report("cp1251→docx", False, str(e))


def test_txt_cp1251_to_txt(files: dict) -> None:
    print("Тест 5: TXT в CP1251 → TXT (перекодировка в UTF-8)")
    try:
        blob, _ = run_convert_job(files["txt_cp1251"], "txt", "e2e cp1251→txt")
    except Exception as e:
        report("cp1251→txt задача", False, str(e))
        return
    try:
        text = blob.decode("utf-8")  # строго UTF-8
        ok, why = has_all_cyrillic(text)
        report("cp1251→txt: результат UTF-8, текст совпадает", ok, why)
    except UnicodeDecodeError as e:
        report("cp1251→txt: результат UTF-8", False, f"НЕ UTF-8 (крякозябры): {e}")


def test_pdf_to_docx(artifacts: dict) -> None:
    print("Тест 6: PDF → DOCX («конвертация назад», pdf2docx)")
    if "pdf" not in artifacts:
        report("pdf→docx", False, "нет входного PDF из теста 2")
        return
    try:
        blob, _ = run_convert_job(artifacts["pdf"], "docx", "e2e pdf→docx")
        text = extract_docx_text(blob)
        ok, why = has_all_cyrillic(text)
        report("pdf→docx: текст восстановлен", ok, why)
    except Exception as e:
        report("pdf→docx", False, str(e))


def test_pdf_to_txt(artifacts: dict) -> None:
    print("Тест 7: PDF → TXT («конвертация назад»)")
    if "pdf" not in artifacts:
        report("pdf→txt", False, "нет входного PDF из теста 2")
        return
    try:
        blob, _ = run_convert_job(artifacts["pdf"], "txt", "e2e pdf→txt")
        text = blob.decode("utf-8")  # строго UTF-8
        ok, why = has_all_cyrillic(text)
        report("pdf→txt: UTF-8, кириллица без крякозябр", ok, why)
    except UnicodeDecodeError as e:
        report("pdf→txt", False, f"НЕ UTF-8 (крякозябры): {e}")
    except Exception as e:
        report("pdf→txt", False, str(e))


def main() -> int:
    print(f"=== E2E конвертера документов VoxPaperAi — движок {BASE} ===")
    print(f"Python: {sys.version.split()[0]}, платформа: {sys.platform}")
    tmp = tempfile.mkdtemp(prefix="vpai_e2e_")
    artifacts = {"tmp": tmp}
    try:
        test_health()
        files = make_fixtures(tmp)
        print(f"Фикстуры в {tmp}: {', '.join(os.path.basename(v) for v in files.values())}")
        test_docx_to_txt(files)
        test_docx_to_pdf(files, artifacts)
        test_txt_cp1251_to_pdf(files)
        test_txt_cp1251_to_docx(files)
        test_txt_cp1251_to_txt(files)
        test_pdf_to_docx(artifacts)
        test_pdf_to_txt(artifacts)
    finally:
        print()
    total = passed + len(failures)
    if failures:
        print(f"ИТОГ: {passed}/{total} ПРОЙДЕНО, ПРОВАЛЕНО: {len(failures)}")
        for name, detail in failures:
            print(f"  ✗ {name}: {detail}")
        return 1
    print(f"ИТОГ: {passed}/{total} ПРОЙДЕНО — конвертер работает, крякозябр нет")
    return 0


if __name__ == "__main__":
    sys.exit(main())
