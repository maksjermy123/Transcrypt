import os
import sys
import tempfile
import unittest
from unittest import mock

import docx


sys.path.append(os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

from src.core.translator import Translator


class TestTranslatorLayout(unittest.TestCase):
    def test_docx_layout_translation_saves_translated_docx(self):
        translator = Translator()

        with tempfile.TemporaryDirectory() as tmpdir:
            input_path = os.path.join(tmpdir, "input.docx")
            output_path = os.path.join(tmpdir, "output.docx")

            document = docx.Document()
            document.add_paragraph("Первый абзац")
            table = document.add_table(rows=1, cols=1)
            table.rows[0].cells[0].paragraphs[0].add_run("Ячейка")
            document.save(input_path)

            def fake_translate_batch(texts, **kwargs):
                return [f"TRANS:{text}" for text in texts]

            with mock.patch.object(translator, "translate_batch", side_effect=fake_translate_batch):
                result = translator.translate_document_with_formatting(
                    input_path,
                    output_path,
                    target_lang="en",
                    provider="Google",
                )

            self.assertEqual(result, output_path)
            self.assertTrue(os.path.exists(output_path))

            translated_doc = docx.Document(output_path)
            texts = [p.text for p in translated_doc.paragraphs if p.text.strip()]
            table_text = translated_doc.tables[0].rows[0].cells[0].paragraphs[0].text

            self.assertIn("TRANS:Первый абзац", texts)
            self.assertEqual(table_text, "TRANS:Ячейка")


if __name__ == "__main__":
    unittest.main()
